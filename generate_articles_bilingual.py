from pathlib import Path
from docx import Document
import html, re, json, time, random
import requests


SOURCE = Path("article_source")
OUTPUT_ZH = Path("articles/zh")
OUTPUT_EN = Path("articles/en")
CACHE_FILE = Path("translate_cache.json")

MAX_CHARS = 4000

SEARCH_ZH = []
SEARCH_EN = []

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")


# ==================== 翻译：必应 ====================

def _keep_plus(src, out):
    """Keep the author's own 18 / 18+ choice.

    The translator silently drops the trailing '+', so every '18' token in the
    output is matched positionally against the source: source '18+' -> '18+',
    source '18' -> '18'. If the two sides do not line up, leave it untouched.
    """
    if not out or "18" not in src:
        return out

    src_nums = re.findall(r"18\+?", src)
    out_hits = list(re.finditer(r"18\+?", out))
    if not src_nums or len(src_nums) != len(out_hits):
        return out

    pieces = []
    last = 0
    for num, m in zip(src_nums, out_hits):
        pieces.append(out[last:m.start()])
        pieces.append("18+" if num.endswith("+") else "18")
        last = m.end()
    pieces.append(out[last:])
    return "".join(pieces)


_NAME_FIXES = [
    (r"[Tt]riangular\s+runes?", "Deltarune"),
    (r"\u042f", "R18"),   # 作者用俄语字母 Я 标注 18+
    (r"\bSpomton\b", "Spamton"),   # 作者笔误
]


def _fix_terms(out):
    """Normalise fan-work proper nouns to their official spelling."""
    if not out:
        return out
    for pat, repl in _NAME_FIXES:
        out = re.sub(pat, repl, out)
    return out


# ==================== 英文白名单 ====================
# 这些词不发给翻译服务，直接原样保留（避免被合并、拼错或吞掉）。
WHITELIST = [
    "MR.(ANT)TENNA", "MR.(ANTON)TENNA", "MR.(ANT) TENNA",
    "Spamton G. Spamton", "Spamton NEO", "Spamton Neo",
    "shadow guy", "Shadow Guy", "Spamton.zip",
    "Swapspt", "Swapspamton", "Defernull", "Deltarune",
    "Pippins", "Jongler", "Addison", "Zapper", "Battat",
    "Tenna", "Spamton", "Mike", "Pluey", "Ggdw", "Neo",
    "tenna", "spamton", "battat", "mike", "pippins",
    "addison", "zapper", "jongler", "pluey", "neo",
]

# 词边界可以是任何非字母数字字符
_WL_RE = re.compile(
    r"(?<![A-Za-z0-9])(" +
    "|".join(re.escape(w) for w in sorted(WHITELIST, key=len, reverse=True)) +
    r")(?![A-Za-z0-9])"
)


def _split_keep(text):
    """Split text into (chunk, is_whitelisted) pairs."""
    if not text:
        return [(text, False)]
    out = []
    last = 0
    for m in _WL_RE.finditer(text):
        if m.start() > last:
            out.append((text[last:m.start()], False))
        out.append((m.group(0), True))
        last = m.end()
    if last < len(text):
        out.append((text[last:], False))
    return out


_TOKEN_TPL = "XQTOKEN%dZ"
_TOKEN_RE = re.compile(r"[Xx][Qq][Tt][Oo][Kk][Ee][Nn]\d+[Zz]?")


def _mask_names(text):
    """Replace whitelisted names with placeholders; return text + mapping."""
    store = {}
    parts = _split_keep(text)
    if not any(keep for _, keep in parts):
        return text, store

    out = []
    for chunk, keep in parts:
        if keep:
            tok = _TOKEN_TPL % len(store)
            store[tok] = chunk
            out.append(tok)
        else:
            out.append(chunk)
    return "".join(out), store


def _unmask(out, store):
    """Restore placeholders; tolerate a dropped trailing Z."""
    if not out or not store:
        return out
    for tok, orig in store.items():
        variants = (tok, tok[:-1]) if tok.endswith("Z") else (tok,)
        for v in variants:
            out = re.sub(re.escape(v) + r"(?![A-Za-z0-9])",
                         lambda _m, o=orig: o, out, flags=re.I)
    return _TOKEN_RE.sub("", out)


def _translate_with_whitelist(translator, text):
    """Translate text with whitelisted names preserved.

    Fast path: mask names, send one request. If any name fails to come back,
    redo that text by translating only the non-name spans (several requests,
    but guaranteed correct).
    """
    if not text:
        return text

    masked, store = _mask_names(text)
    if not store:
        return translator(text)

    out = translator(masked)
    if out:
        restored = _unmask(out, store)
        if all(orig in restored for orig in store.values()):
            return restored
    # Fallback: strict, name spans never sent to the translator
    result = []
    for chunk, keep in _split_keep(text):
        if keep:
            result.append(chunk)
        elif chunk.strip():
            result.append(translator(chunk) or chunk)
        else:
            result.append(chunk)
    return "".join(result)


class Bing:
    def __init__(self):
        self.session = requests.Session()
        self.session.trust_env = False
        self.session.proxies = {}
        self.headers = {"User-Agent": UA, "Referer": "https://cn.bing.com/translator"}
        self.ig = ""
        self.iid = "translator.5024"
        self.key = ""
        self.token = ""
        self.refresh()

    def refresh(self):
        for _ in range(3):
            try:
                r = self.session.get("https://cn.bing.com/translator",
                                     headers=self.headers, timeout=20,
                                     proxies={"http": None, "https": None})
                ig = re.search(r'IG:"([A-Za-z0-9]+)"', r.text)
                k = re.search(r'params_AbusePreventionHelper\s*=\s*\[(\d+),"([^"]+)"', r.text)
                if ig and k:
                    self.ig, self.iid, self.key, self.token = ig.group(1), "translator.5024", k.group(1), k.group(2)
                    return True
            except Exception:
                pass
            time.sleep(1.5)
        return False

    def _post(self, text):
        params = {"isVertical": "1", "IG": self.ig, "IID": self.iid}
        data = [("fromLang", "zh-Hans"), ("text", text), ("to", "en"),
                ("token", self.token), ("key", self.key)]
        r = self.session.post("https://cn.bing.com/ttranslatev3", params=params,
                              data=data, headers=self.headers, timeout=45,
                              proxies={"http": None, "https": None})
        if r.status_code != 200:
            return None
        try:
            payload = r.json()
        except Exception:
            return None
        if not isinstance(payload, list) or not payload:
            return None
        part = payload[0].get("translations", [{}])[0].get("text")
        return part

    def translate(self, text, retry=4):
        """Translate, keeping whitelisted English names verbatim."""
        for attempt in range(retry):
            try:
                out = _translate_with_whitelist(self._post, text)
            except Exception:
                out = None
            if out:
                return _fix_terms(_keep_plus(text, out))
            time.sleep(1.0 + attempt)
            self.refresh()
        return None

    def translate_paragraphs(self, paras):
        """把段落用换行拼成 <4000 字符的块，一次请求翻译，再拆回。"""
        result = []
        i = 0
        while i < len(paras):
            chunk = []
            size = 0
            while i < len(paras) and size + len(paras[i]) + 1 <= MAX_CHARS:
                chunk.append(paras[i])
                size += len(paras[i]) + 1
                i += 1
            if not chunk:
                chunk = [paras[i]]
                i += 1

            joined = "\n".join(chunk)
            out = self.translate(joined)
            if out is not None:
                out = _fix_terms(_keep_plus(joined, out))
            if out is None:
                result.extend(chunk)
                continue

            lines = [s.strip() for s in out.split("\n")]
            lines = [s for s in lines if s != ""]
            if len(lines) == len(chunk):
                result.extend(lines)
            else:
                for p in chunk:
                    single = self.translate(p)
                    result.append(single if single else p)
                    time.sleep(0.15)
            time.sleep(0.25)
        return result


# ==================== UI 文案 ====================
UI = {
    "zh": {"html_lang": "zh-CN", "back_prev": "← 返回上一级", "back_home": "← 返回首页",
           "comment": "留言", "comment_ph": "写下你的留言", "send": "发送",
           "music_cur": "当前音乐：", "music_none": "未播放", "folder": "📁", "article": "📄",
           "dirs": "目录", "arts": "文章", "lang_label": "EN", "lang_title": "Switch to English",
           "search_ph": "搜索文章...",
           "no_content": "暂无内容", "unreadable": "文件无法读取", "no_result": "没有找到相关内容"},
    "en": {"html_lang": "en", "back_prev": "← Back", "back_home": "← Home",
           "comment": "Comments", "comment_ph": "Leave a comment", "send": "Send",
           "music_cur": "Now playing: ", "music_none": "Not playing", "folder": "📁", "article": "📄",
           "dirs": "Contents", "arts": "Articles", "lang_label": "中文", "lang_title": "切换到中文",
           "search_ph": "Search articles...",
           "no_content": "No content", "unreadable": "File cannot be read", "no_result": "No results found"},
}


# ==================== 模板 ====================
MUSIC = """<audio id="bgm" loop></audio>

<div class="music-box">
<div class="music-title"><span class="music-icon">🎵</span> {cur}<span id="musicName">{none}</span></div>
<div class="music-line"></div>
<div class="music-controls">
<button class="music-btn small" onclick="prevMusic()">⏮</button>
<button id="playBtn" class="music-btn play" onclick="toggleMusic()">▶</button>
<button class="music-btn small" onclick="nextMusic()">⏭</button>
<select id="songSelect"></select>
</div>
</div>"""

SEARCH_BOX = """<div class="search-box">
<input id="site-search" placeholder="{ph}">
<div id="searchResult"></div>
</div>"""

ARTICLE_TPL = """<!DOCTYPE html>
<html lang="{html_lang}">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<link rel="stylesheet" href="{css}">
</head>
<body>

<div id="lang-btn" class="lang-switch" title="{lang_title}" onclick="switchLang()">{lang_label}</div>
<button id="theme-btn">🌙</button>

{search_box}

<div class="container">
<div class="project-page">

<div class="nav-bar">
<a class="back-btn" href="index.html">{back_prev}</a>
<a class="back-btn" href="{home}">{back_home}</a>
</div>

<h1>{title}</h1>

<div class="article-content">
{content}
</div>

<div class="nav-bar">
{article_nav}
</div>

<div class="comment-box">
<h3>{comment}</h3>
<textarea id="message" placeholder="{comment_ph}"></textarea>
<button onclick="sendComment()">{send}</button>
</div>

</div>
</div>

{music}

<script src="{theme}"></script>
<script src="{script}"></script>
<script src="{search_js}"></script>
{lang_script}

</body>
</html>"""

INDEX_TPL = """<!DOCTYPE html>
<html lang="{html_lang}">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<link rel="stylesheet" href="{css}">
</head>
<body>

<div id="lang-btn" class="lang-switch" title="{lang_title}" onclick="switchLang()">{lang_label}</div>
<button id="theme-btn">🌙</button>

{search_box}

<div class="container">
<div class="project-page">

<div class="nav-bar">
{nav}
</div>

<h1>{title}</h1>

{folders}

{articles}

</div>
</div>

{music}

<script src="{theme}"></script>
<script src="{script}"></script>
<script src="{search_js}"></script>
{lang_script}

</body>
</html>"""

LANG_SCRIPT = """<script>
(function () {
  var lang = localStorage.getItem("lang") || "zh";
  document.documentElement.lang = lang === "en" ? "en" : "zh-CN";
  window.switchLang = function () {
    var cur = localStorage.getItem("lang") || "zh";
    var next = cur === "zh" ? "en" : "zh";
    localStorage.setItem("lang", next);
    var p = window.location.pathname;
    window.location.href = p.replace("/articles/" + cur + "/", "/articles/" + next + "/");
  };
  if (lang === "en" && window.location.pathname.indexOf("/articles/zh/") !== -1) {
    window.location.pathname = window.location.pathname.replace("/articles/zh/", "/articles/en/");
  } else if (lang === "zh" && window.location.pathname.indexOf("/articles/en/") !== -1) {
    window.location.pathname = window.location.pathname.replace("/articles/en/", "/articles/zh/");
  }
})();
</script>"""


# ==================== 缓存 ====================
CACHE = {}

def load_cache():
    global CACHE
    if CACHE_FILE.exists():
        try:
            CACHE = json.loads(CACHE_FILE.read_text(encoding="utf-8"))
        except Exception:
            CACHE = {}
    # 清掉无效条目（译文和原文完全相同 = 当初翻译失败）
    bad = [k for k, v in CACHE.items() if k == v]
    for k in bad:
        del CACHE[k]
    print("缓存有效条目:", len(CACHE), "（丢弃无效", len(bad), "条）")

def save_cache():
    CACHE_FILE.write_text(json.dumps(CACHE, ensure_ascii=False, indent=2), encoding="utf-8")


# ==================== 工具 ====================
def natural_sort(items):
    return sorted(items, key=lambda x: [int(s) if s.isdigit() else s.lower()
                                         for s in re.split(r"(\d+)", x.name)])

def read_paras(f, ui):
    try:
        doc = Document(f)
    except Exception:
        return [ui["unreadable"]]
    ps = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return ps or [ui["no_content"]]

def to_html(paras):
    return "\n".join("<p>%s</p>" % html.escape(p) for p in paras)

def strip_tags(s):
    return re.sub("<.*?>", "", s)

def depth(out, base):
    return len(out.relative_to(base).parts)

def up(out, base, extra=0):
    return "../" * (depth(out, base) + extra)

def rel_depth(out, base, extra=0):
    # 相对于 articles/<lang>/ 的深度（用于 style.css 等根文件）
    return depth(out, base) + 1 + extra


def asset(out, base, name):
    return "../" * (depth(out, base) + 1) + name


def to_root(out, base):
    return "../" * (depth(out, base) + 1) + "index.html"


# ==================== 生成 ====================
def build_article(src, oz, oe, uz, ue, tz, te, nav_zh, nav_en):
    paras = read_paras(src, uz)
    zh_html_content = to_html(paras)

    # 中文页面（无条件生成）
    music = MUSIC.format(cur=uz["music_cur"], none=uz["music_none"])
    oz.write_text(ARTICLE_TPL.format(
        html_lang=uz["html_lang"], title=html.escape(tz), content=zh_html_content,
        css=asset(oz, OUTPUT_ZH, "style.css"), theme=asset(oz, OUTPUT_ZH, "theme.js"),
        script=asset(oz, OUTPUT_ZH, "script.js"), home=to_root(oz, OUTPUT_ZH),
        article_nav="".join(nav_zh), music=music, lang_script=LANG_SCRIPT,
        search_box=SEARCH_BOX.format(ph=uz["search_ph"]),
        search_js=up(oz, OUTPUT_ZH, 1) + "search.js",
        back_prev=uz["back_prev"], back_home=uz["back_home"], comment=uz["comment"],
        comment_ph=uz["comment_ph"], send=uz["send"], lang_label=uz["lang_label"],
        lang_title=uz["lang_title"]), encoding="utf-8")

    SEARCH_ZH.append({"title": tz, "path": str(oz).replace("\\", "/"),
                      "content": strip_tags(zh_html_content)})

    # 英文页面（翻译，失败则保留原文）
    cached_paras = []
    for p in paras:
        if p in CACHE:
            cached_paras.append(CACHE[p])
        else:
            cached_paras.append(None)

    missing = [p for p, c in zip(paras, cached_paras) if c is None]
    if missing:
        translated = BING.translate_paragraphs(missing)
    else:
        translated = []
    it = iter(translated)
    en_paras = [c if c is not None else next(it) for c in cached_paras]

    for p, e in zip(paras, en_paras):
        CACHE[p] = _fix_terms(_keep_plus(p, e))

    en_title = CACHE.get(tz)
    if not en_title:
        en_title = BING.translate(tz) or tz
    en_title = _fix_terms(_keep_plus(tz, en_title))
    CACHE[tz] = en_title

    en_paras = [_fix_terms(_keep_plus(src_p, e)) for src_p, e in zip(paras, en_paras)]
    en_html_content = to_html(en_paras)
    music_en = MUSIC.format(cur=ue["music_cur"], none=ue["music_none"])
    oe.write_text(ARTICLE_TPL.format(
        html_lang=ue["html_lang"], title=html.escape(en_title), content=en_html_content,
        css=asset(oe, OUTPUT_EN, "style.css"), theme=asset(oe, OUTPUT_EN, "theme.js"),
        script=asset(oe, OUTPUT_EN, "script.js"), home=to_root(oe, OUTPUT_EN),
        article_nav="".join(nav_en), music=music_en, lang_script=LANG_SCRIPT,
        search_box=SEARCH_BOX.format(ph=ue["search_ph"]),
        search_js=up(oe, OUTPUT_EN, 1) + "search.js",
        back_prev=ue["back_prev"], back_home=ue["back_home"], comment=ue["comment"],
        comment_ph=ue["comment_ph"], send=ue["send"], lang_label=ue["lang_label"],
        lang_title=ue["lang_title"]), encoding="utf-8")

    SEARCH_EN.append({"title": en_title, "path": str(oe).replace("\\", "/"),
                      "content": strip_tags(en_html_content)})


def build_index(folder, oz, oe, uz, ue):
    folders_zh, folders_en, arts_zh, arts_en = [], [], [], []

    for item in natural_sort(list(folder.iterdir())):
        if item.is_dir():
            name_en = CACHE.get(item.name) or BING.translate(item.name) or item.name
            CACHE[item.name] = name_en
            folders_zh.append('<a class="article-card" href="%s/index.html"><h3>%s %s</h3></a>'
                              % (item.name, uz["folder"], html.escape(item.name)))
            folders_en.append('<a class="article-card" href="%s/index.html"><h3>%s %s</h3></a>'
                              % (item.name, ue["folder"], html.escape(name_en)))
            SEARCH_ZH.append({"title": item.name,
                              "path": str(oz.parent / item.name / "index.html").replace("\\", "/"),
                              "content": item.name})
            SEARCH_EN.append({"title": name_en,
                              "path": str(oe.parent / item.name / "index.html").replace("\\", "/"),
                              "content": name_en})
        elif item.suffix.lower() == ".docx" and not item.name.startswith("~$"):
            t = item.stem
            t_en = CACHE.get(t) or BING.translate(t) or t
            CACHE[t] = t_en
            arts_zh.append('<a class="article-card" href="%s.html"><h3>%s %s</h3></a>'
                           % (t, uz["article"], html.escape(t)))
            arts_en.append('<a class="article-card" href="%s.html"><h3>%s %s</h3></a>'
                           % (t, ue["article"], html.escape(t_en)))

    dir_zh = '<h2>%s</h2>\n<div class="article-grid">%s</div>' % (uz["dirs"], "".join(folders_zh)) if folders_zh else ""
    dir_en = '<h2>%s</h2>\n<div class="article-grid">%s</div>' % (ue["dirs"], "".join(folders_en)) if folders_en else ""
    art_zh = '<h2>%s</h2>\n<div class="article-grid">%s</div>' % (uz["arts"], "".join(arts_zh)) if arts_zh else ""
    art_en = '<h2>%s</h2>\n<div class="article-grid">%s</div>' % (ue["arts"], "".join(arts_en)) if arts_en else ""

    is_entry = len(oz.parent.relative_to(OUTPUT_ZH).parts) <= 1
    if is_entry:
        nav_zh = '<a class="back-btn" href="%s">%s</a>' % (to_root(oz, OUTPUT_ZH), uz["back_home"])
        nav_en = '<a class="back-btn" href="%s">%s</a>' % (to_root(oe, OUTPUT_EN), ue["back_home"])
    else:
        nav_zh = ('<a class="back-btn" href="../index.html">%s</a>\n'
                  '<a class="back-btn" href="%s">%s</a>') % (uz["back_prev"], to_root(oz, OUTPUT_ZH), uz["back_home"])
        nav_en = ('<a class="back-btn" href="../index.html">%s</a>\n'
                  '<a class="back-btn" href="%s">%s</a>') % (ue["back_prev"], to_root(oe, OUTPUT_EN), ue["back_home"])

    t_title = folder.name
    t_title_en = CACHE.get(t_title) or BING.translate(t_title) or t_title
    CACHE[t_title] = t_title_en

    music_zh = MUSIC.format(cur=uz["music_cur"], none=uz["music_none"])
    music_en = MUSIC.format(cur=ue["music_cur"], none=ue["music_none"])

    oz.write_text(INDEX_TPL.format(
        html_lang=uz["html_lang"], title=html.escape(t_title), folders=dir_zh, articles=art_zh,
        nav=nav_zh, music=music_zh, css=asset(oz, OUTPUT_ZH, "style.css"),
        theme=asset(oz, OUTPUT_ZH, "theme.js"), script=asset(oz, OUTPUT_ZH, "script.js"),
        search_js=up(oz, OUTPUT_ZH, 1) + "search.js", lang_script=LANG_SCRIPT,
        search_box=SEARCH_BOX.format(ph=uz["search_ph"]),
        lang_label=uz["lang_label"], lang_title=uz["lang_title"]), encoding="utf-8")

    oe.write_text(INDEX_TPL.format(
        html_lang=ue["html_lang"], title=html.escape(t_title_en), folders=dir_en, articles=art_en,
        nav=nav_en, music=music_en, css=asset(oe, OUTPUT_EN, "style.css"),
        theme=asset(oe, OUTPUT_EN, "theme.js"), script=asset(oe, OUTPUT_EN, "script.js"),
        search_js=up(oe, OUTPUT_EN, 1) + "search.js", lang_script=LANG_SCRIPT,
        search_box=SEARCH_BOX.format(ph=ue["search_ph"]),
        lang_label=ue["lang_label"], lang_title=ue["lang_title"]), encoding="utf-8")


def walk(src, oz, oe, uz, ue, counter):
    oz.mkdir(parents=True, exist_ok=True)
    oe.mkdir(parents=True, exist_ok=True)

    docs = [x for x in natural_sort(list(src.iterdir()))
            if x.suffix.lower() == ".docx" and not x.name.startswith("~$")]

    for i, d in enumerate(docs):
        counter[0] += 1
        tz = d.stem
        te = CACHE.get(tz)
        if not te:
            te = BING.translate(tz) or tz
            CACHE[tz] = te

        nav_zh, nav_en = [], []
        if i > 0:
            p = docs[i - 1]
            p_en = CACHE.get(p.stem) or BING.translate(p.stem) or p.stem
            CACHE[p.stem] = p_en
            nav_zh.append('<a class="back-btn" href="%s.html">← %s</a>' % (p.stem, html.escape(p.stem)))
            nav_en.append('<a class="back-btn" href="%s.html">← %s</a>' % (p.stem, html.escape(p_en)))
        if i < len(docs) - 1:
            n = docs[i + 1]
            n_en = CACHE.get(n.stem) or BING.translate(n.stem) or n.stem
            CACHE[n.stem] = n_en
            nav_zh.append('<a class="back-btn" href="%s.html">%s →</a>' % (n.stem, html.escape(n.stem)))
            nav_en.append('<a class="back-btn" href="%s.html">%s →</a>' % (n.stem, html.escape(n_en)))

        print("  [%d/%d] %s" % (counter[0], counter[1], d.name))
        try:
            build_article(d, oz / (d.stem + ".html"), oe / (d.stem + ".html"),
                          uz, ue, tz, te, nav_zh, nav_en)
        except Exception as e:
            print("    !! 跳过: %s" % e)
            time.sleep(2)
            BING.refresh()

        if counter[0] % 10 == 0:
            save_cache()

    for item in natural_sort(list(src.iterdir())):
        if item.is_dir():
            walk(item, oz / item.name, oe / item.name, uz, ue, counter)

    build_index(src, oz / "index.html", oe / "index.html", uz, ue)
    save_cache()


def main():
    global BING
    if not SOURCE.exists():
        print("未找到 article_source")
        return

    total = sum(1 for _ in SOURCE.rglob("*.docx") if not _.name.startswith("~$"))
    print("共 %d 篇 docx" % total)
    load_cache()
    BING = Bing()
    print("必应翻译会话:", "OK" if BING.token else "FAILED")

    t0 = time.time()
    walk(SOURCE, OUTPUT_ZH, OUTPUT_EN, UI["zh"], UI["en"], [0, total])

    (OUTPUT_ZH / "search.json").write_text(
        json.dumps(SEARCH_ZH, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUT_EN / "search.json").write_text(
        json.dumps(SEARCH_EN, ensure_ascii=False, indent=2), encoding="utf-8")

    # 根目录索引：供首页 / 周边页的搜索框按语言取用
    (Path("search_zh.json")).write_text(
        json.dumps(SEARCH_ZH, ensure_ascii=False, indent=2), encoding="utf-8")
    (Path("search_en.json")).write_text(
        json.dumps(SEARCH_EN, ensure_ascii=False, indent=2), encoding="utf-8")
    print("  已写出 search_zh.json / search_en.json")

    print("\n完成，耗时 %.0f 秒" % (time.time() - t0))
    print("  中文: %d 条" % len(SEARCH_ZH))
    print("  英文: %d 条" % len(SEARCH_EN))
    print("  缓存: %d 条" % len(CACHE))


if __name__ == "__main__":
    main()